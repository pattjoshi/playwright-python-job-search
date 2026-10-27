"""Scraped job data is untrusted: it can't run formulas in Excel."""

import csv

from openpyxl import load_workbook

from job_search.models import Job
from job_search.services.reporting.excel_csv import as_text, write_reports


def test_spreadsheet_cells_never_become_formulas(tmp_path):
    job = Job("naukri", "1", '=HYPERLINK("http://evil.example","Click")', "+Acme", "@Pune", "file:///C:/Windows/calc.exe",
              match_reason="-2+3", score=80)
    xlsx, csv_path = write_reports([job], tmp_path, "jobs")
    row = list(load_workbook(xlsx).active.iter_rows(min_row=2, values_only=True))[0]
    assert row[1:4] == ("'=HYPERLINK(\"http://evil.example\",\"Click\")", "'+Acme", "'@Pune")
    link = load_workbook(xlsx).active.cell(row=2, column=11)
    assert link.hyperlink is None  # only http(s) links are clickable
    with csv_path.open(encoding="utf-8-sig") as handle:
        assert list(csv.reader(handle))[1][1].startswith("'=")
    assert as_text("AI Engineer") == "AI Engineer" and as_text(85) == 85

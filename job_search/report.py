"""Write ranked jobs to Excel and CSV."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from job_search.models import Job

COLUMNS = [
    ("Score", 8, lambda j: j.score),
    ("Title", 40, lambda j: j.title),
    ("Company", 28, lambda j: j.company),
    ("Location", 28, lambda j: j.location),
    ("Posted", 16, lambda j: j.posted_text),
    ("Experience", 12, lambda j: j.experience),
    ("Source", 10, lambda j: j.source),
    ("Matching skills", 40, lambda j: ", ".join(j.matched_skills)),
    ("Missing skills", 30, lambda j: ", ".join(j.missing_skills)),
    ("Why it matches", 70, lambda j: j.match_reason),
    ("Link", 50, lambda j: j.url),
]


def write_reports(jobs: list[Job], output_dir: Path, stem: str) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = output_dir / f"{stem}.xlsx"
    csv_path = output_dir / f"{stem}.csv"
    _write_xlsx(jobs, xlsx_path)
    _write_csv(jobs, csv_path)
    return xlsx_path, csv_path


def _write_xlsx(jobs: list[Job], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jobs"
    sheet.append([name for name, _, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (_, width, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    link_column = len(COLUMNS)
    for job in jobs:
        sheet.append([getter(job) for _, _, getter in COLUMNS])
        link_cell = sheet.cell(row=sheet.max_row, column=link_column)
        if job.url:
            link_cell.hyperlink = job.url
            link_cell.style = "Hyperlink"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def _write_csv(jobs: list[Job], path: Path) -> None:
    # utf-8-sig so Excel opens non-English characters correctly.
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([name for name, _, _ in COLUMNS])
        for job in jobs:
            writer.writerow(["" if (value := getter(job)) is None else value for _, _, getter in COLUMNS])

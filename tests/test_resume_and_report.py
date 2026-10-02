from pathlib import Path

import pytest
from openpyxl import load_workbook

from job_search.models import Job
from job_search.report import write_reports
from job_search.resume import read_resume_text

SAMPLE_RESUME = Path(__file__).parent / "sample_resume.txt"


def test_read_text_resume():
    assert "Playwright" in read_resume_text(SAMPLE_RESUME)


def test_read_docx_resume(tmp_path):
    import docx

    path = tmp_path / "resume.docx"
    document = docx.Document()
    document.add_paragraph("Python developer with Django experience")
    document.save(path)

    assert "Django" in read_resume_text(path)


def test_resume_errors(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_resume_text(tmp_path / "missing.pdf")
    (tmp_path / "resume.rtf").write_text("x")
    with pytest.raises(ValueError, match="Unsupported"):
        read_resume_text(tmp_path / "resume.rtf")
    (tmp_path / "empty.txt").write_text("  ")
    with pytest.raises(ValueError, match="No text"):
        read_resume_text(tmp_path / "empty.txt")


def test_write_reports(tmp_path):
    jobs = [
        Job("linkedin", "1", "SDET", "Acme", "Pune", "https://example.com/1", "2 hours ago", score=88, match_reason="Fit"),
        Job("linkedin", "2", "QA", "Globex", "Delhi", "https://example.com/2"),
    ]

    xlsx_path, csv_path = write_reports(jobs, tmp_path / "out", "jobs_test")

    sheet = load_workbook(xlsx_path).active
    assert sheet["A1"].value == "Score"
    assert sheet["A2"].value == 88
    assert sheet["J2"].hyperlink.target == "https://example.com/1"
    lines = csv_path.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].startswith("Score,Title")
    assert lines[2].startswith(",QA,Globex")
